"""Excel/CSV export. The Summary sheet is formula-driven off the Reviews sheet, so edits to the
data, theme keywords or month starts recalculate live (Excel / LibreOffice / Google Sheets)."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .normalize import sentiment_of

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SEC_BORDER = Border(bottom=Side(style="medium", color="1F3A5F"))
BLUE = "0000FF"  # editable inputs
DEFAULT_THEME_KEYWORDS = ["bread", "price", "expensive", "wait", "slow", "staff", "rude", "clean", "dirty",
                          "fresh", "portion", "wrong", "delivery", "app", "cookie"]


def get_theme_keywords(brand: str = "Subway") -> list[str]:
    from .menu import load_menu
    base = ["price", "expensive", "wait", "slow", "staff", "rude", "clean", "dirty",
            "fresh", "portion", "wrong", "delivery", "app"]
    menu = load_menu(brand)
    specific = []
    for it in list(menu.keys())[:5]:
        kw = it.lower().split()[0].replace(".", "").strip()
        if kw and len(kw) > 2 and kw not in base and kw not in specific:
            specific.append(kw)
    return specific + base

REVIEW_COLS = ["Source", "Place", "Address", "Rating", "Date", "Review Text", "Likes",
               "Owner Response", "Author", "Sentiment", "Review ID", "Place URL"]
CSV_COLS = ["source", "place_name", "address", "rating", "review_date", "text", "likes",
            "owner_response", "author", "sentiment", "review_id", "place_url", "place_id"]


def _clean(v, limit: int = 32000):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:limit]
    return v


def put(ws, row, col, value, *, bold=False, color=None, fmt=None, text=False, size=10, fill=None,
        wrap=False, align=None):
    c = ws.cell(row=row, column=col, value=_clean(value))
    if text and isinstance(c.value, str):
        c.data_type = "s"  # review text starting with "=" must never become a formula
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top")
    return c


def _header(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        put(ws, row, start_col + i, label, bold=True, color="FFFFFF", fill=HEAD_FILL, align="center")


def _section(ws, row, title, width=4):
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).border = SEC_BORDER
    put(ws, row, 1, title, bold=True, size=11).border = SEC_BORDER


def _d(s):
    return datetime.strptime(s, "%Y-%m-%d").date() if s else None


def build_workbook(path: Path, run: dict, reviews: list[dict], places: list[dict], jobs: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rv = wb.create_sheet("Reviews")
    pl = wb.create_sheet("Places")
    info = wb.create_sheet("Run Info")

    # ------------------------------------------------------------ Reviews
    _header(rv, 1, REVIEW_COLS)
    for i, r in enumerate(reviews, start=2):
        put(rv, i, 1, r["source"], text=True)
        put(rv, i, 2, r["place_name"], text=True)
        put(rv, i, 3, r["address"], text=True)
        put(rv, i, 4, r["rating"], fmt="0.0")
        put(rv, i, 5, _d(r["review_date"]), fmt="yyyy-mm-dd")
        put(rv, i, 6, r["text"], text=True)
        put(rv, i, 7, r["likes"])
        put(rv, i, 8, r["owner_response"], text=True)
        put(rv, i, 9, r["author"], text=True)
        put(rv, i, 10, f'=IF(D{i}="","",IF(D{i}>=4,"Positive",IF(D{i}>=3,"Neutral","Negative")))')
        put(rv, i, 11, r["review_id"], text=True)
        put(rv, i, 12, r["place_url"], text=True)
    for col, w in zip("ABCDEFGHIJKL", [13, 22, 30, 8, 12, 80, 7, 30, 18, 11, 22, 40]):
        rv.column_dimensions[col].width = w
    rv.freeze_panes = "A2"
    rv.auto_filter.ref = f"A1:L{max(len(reviews) + 1, 2)}"

    # ------------------------------------------------------------ Places
    _header(pl, 1, ["Source", "Place", "Address", "Google Rating", "Google Review Count", "Lat", "Lon", "URL"])
    for i, p in enumerate(places, start=2):
        for col, key in enumerate(["source", "name", "address", "rating", "reviews_count", "lat", "lon", "url"], 1):
            put(pl, i, col, p.get(key), text=True)
    for col, w in zip("ABCDEFGH", [13, 22, 36, 14, 20, 11, 11, 60]):
        pl.column_dimensions[col].width = w
    pl.freeze_panes = "A2"

    # ------------------------------------------------------------ Summary (formulas)
    R = "Reviews!$%s:$%s"
    SRC, RAT, DATE, TEXT, RESP = (R % (c, c) for c in "ADEFH")
    ws.column_dimensions["A"].width = 34
    for col in "BCD":
        ws.column_dimensions[col].width = 16
    put(ws, 1, 1, f'{run["brand"]} - Customer Review Summary', bold=True, size=14)
    put(ws, 2, 1, "Run ID"); put(ws, 2, 2, run["id"], text=True)
    put(ws, 3, 1, "Exported (UTC)"); put(ws, 3, 2, datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"), text=True)
    window = run["params"].get("days_limit")
    put(ws, 4, 1, "Google review window"); put(ws, 4, 2, f"last {window} days" if window else "all history", text=True)

    _section(ws, 6, "Key metrics")
    kpis = [
        ("Total reviews", f"=COUNTA({SRC})-1", "#,##0"),
        ("Average rating (1-5)", f"=IFERROR(AVERAGE({RAT}),0)", "0.00"),
        ("Negative share (rating <= 2)", f'=IFERROR(COUNTIF({RAT},"<=2")/B7,0)', "0.0%"),
        ("Positive share (rating >= 4)", f'=IFERROR(COUNTIF({RAT},">=4")/B7,0)', "0.0%"),
        ("Reviews with an owner response", f'=MAX(COUNTIF({RESP},"?*")-1,0)', "#,##0"),
    ]
    for i, (label, formula, fmt) in enumerate(kpis, start=7):
        put(ws, i, 1, label); put(ws, i, 2, formula, fmt=fmt, bold=True)

    row = 13
    _section(ws, row, "By source"); row += 1
    _header(ws, row, ["Source", "Reviews", "Avg rating", "Negative share"]); row += 1
    for src in sorted({r["source"] for r in reviews}) or ["google_maps"]:
        put(ws, row, 1, src, color=BLUE)
        put(ws, row, 2, f"=COUNTIF({SRC},A{row})", fmt="#,##0")
        put(ws, row, 3, f"=IFERROR(AVERAGEIF({SRC},A{row},{RAT}),0)", fmt="0.00")
        put(ws, row, 4, f'=IFERROR(COUNTIFS({SRC},A{row},{RAT},"<=2")/B{row},0)', fmt="0.0%")
        row += 1

    row += 1
    _section(ws, row, "Rating distribution"); row += 1
    _header(ws, row, ["Stars", "Reviews", "Share"]); dist_head = row; row += 1
    dist_first = row
    for stars in range(1, 6):
        put(ws, row, 1, stars, color=BLUE)
        put(ws, row, 2, f"=COUNTIF({RAT},A{row})", fmt="#,##0")
        put(ws, row, 3, f"=IFERROR(B{row}/$B$7,0)", fmt="0.0%")
        row += 1
    chart = BarChart()
    chart.title, chart.y_axis.title, chart.x_axis.title = "Reviews by star rating", "Reviews", "Stars"
    chart.add_data(Reference(ws, min_col=2, min_row=dist_head, max_row=row - 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=dist_first, max_row=row - 1))
    chart.legend = None
    chart.height, chart.width = 7, 13
    ws.add_chart(chart, "F6")

    dates = sorted(_d(r["review_date"]) for r in reviews if r["review_date"])
    if dates:
        row += 1
        _section(ws, row, "Reviews by month (edit the start dates)"); row += 1
        _header(ws, row, ["Month start", "Reviews", "Avg rating"]); row += 1
        last = dates[-1].replace(day=1)
        months = []
        y, m = last.year, last.month
        for _ in range(6):
            months.append(date(y, m, 1))
            y, m = (y - 1, 12) if m == 1 else (y, m - 1)
        for ms in reversed(months):
            put(ws, row, 1, ms, color=BLUE, fmt="yyyy-mm-dd")
            win = f'{DATE},">="&A{row},{DATE},"<"&EDATE(A{row},1)'
            put(ws, row, 2, f"=COUNTIFS({win})", fmt="#,##0")
            put(ws, row, 3, f"=IFERROR(AVERAGEIFS({RAT},{win}),0)", fmt="0.00")
            row += 1

    row += 1
    _section(ws, row, "Themes (keyword mentions; edit the blue keywords)"); row += 1
    _header(ws, row, ["Keyword", "Mentions", "In negative reviews", "Negative share"]); row += 1
    theme_kws = get_theme_keywords(run.get("brand", "Subway"))
    for kw in theme_kws:
        put(ws, row, 1, kw, color=BLUE)
        put(ws, row, 2, f'=COUNTIF({TEXT},"*"&A{row}&"*")', fmt="#,##0")
        put(ws, row, 3, f'=COUNTIFS({TEXT},"*"&A{row}&"*",{RAT},"<=2")', fmt="#,##0")
        put(ws, row, 4, f"=IFERROR(C{row}/B{row},0)", fmt="0.0%")
        row += 1
    put(ws, row + 1, 1, "Blue = editable input. Theme counts are case-insensitive substring matches "
                        "on review text, not NLP.", size=9)

    # ------------------------------------------------------------ Run Info (static metadata)
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 110
    put(info, 1, 1, "Run parameters and notes", bold=True, size=12)
    p = run["params"]
    rows = [("Brand", run["brand"]), ("Status", run["status"]), ("Sources", ", ".join(p["sources"])),
            ("max_places", p["max_places"]), ("days_limit", p.get("days_limit") or "all"),
            ("Reviews exported (static count)", len(reviews)), ("Places exported (static count)", len(places))]
    for i, (k, v) in enumerate(rows, start=3):
        put(info, i, 1, k, bold=True); put(info, i, 2, v, text=True)
    r0 = 3 + len(rows) + 1
    put(info, r0, 1, "Jobs", bold=True, size=11)
    _header(info, r0 + 1, ["Source / kind", "Result"])
    for i, j in enumerate(jobs, start=r0 + 2):
        put(info, i, 1, f'{j["source"]} / {j["kind"]}'); put(info, i, 2, f'{j["status"]}: {j.get("message") or ""}', text=True)
    n0 = r0 + len(jobs) + 3
    notes = [
        "Sentiment rule: rating >= 4 Positive, >= 3 Neutral, otherwise Negative (Reviews!J).",
        "Duplicate reviews are removed per (source, review_id). Review text is truncated at 32,000 characters.",
        "Reviews from Google Maps are the verified schema; Tripadvisor/Yelp field mapping is tolerant, check Reviews!B:F on first live run.",
        "Source data is public customer content collected via Bright Data; review authors' names are personal data under GDPR/CCPA.",
    ]
    for i, n in enumerate(notes, start=n0):
        put(info, i, 1, "Note"); put(info, i, 2, n, text=True)

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def csv_bytes(reviews: list[dict]) -> bytes:
    """CSV with formula-injection guard (cells starting = + - @ are prefixed with a quote)."""
    def safe(v):
        if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
            return "'" + v
        return v
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CSV_COLS)
    for r in reviews:
        row = {**r, "sentiment": sentiment_of(r["rating"])}
        w.writerow([safe(row.get(c)) for c in CSV_COLS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


def build_insights_workbook(path: Path, run: dict, items: list[dict], matrix: dict, join: dict | None) -> Path:
    """Static analytic values (computed in Python); the formula-driven workbook is export.xlsx."""
    from openpyxl.formatting.rule import ColorScaleRule
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    put(ws, 1, 1, f'{run["brand"]} - Sentiment by menu item (run {run["id"]})', bold=True, size=13)
    _header(ws, 3, ["Item", "Mentions", "Reviews", "Avg sentiment (-1..1)", "Negative share", "Top aspects"])
    for i, it in enumerate(items, start=4):
        put(ws, i, 1, it["item"], text=True); put(ws, i, 2, it["mentions"]); put(ws, i, 3, it["reviews"])
        put(ws, i, 4, it["avg_sentiment"], fmt="0.00"); put(ws, i, 5, it["negative_share"], fmt="0.0%")
        put(ws, i, 6, "; ".join(f'{a["aspect"]} ({a["mentions"]}, {a["avg_sentiment"]})' for a in it["top_aspects"]), text=True)
    for col, w in zip("ABCDEF", [28, 10, 10, 20, 15, 70]):
        ws.column_dimensions[col].width = w
    if items:
        ws.conditional_formatting.add(f"D4:D{3 + len(items)}", ColorScaleRule(
            start_type="num", start_value=-1, start_color="F8696B", mid_type="num", mid_value=0,
            mid_color="FFFFFF", end_type="num", end_value=1, end_color="63BE7B"))
    ws.freeze_panes = "A4"

    mx = wb.create_sheet("Aspect Matrix")
    aspects = sorted({a for d in matrix.values() for a in d})
    put(mx, 1, 1, "Average sentiment by item x aspect (cell shows avg; mentions in the second block)", bold=True)
    _header(mx, 3, ["Item"] + aspects)
    for i, item in enumerate(sorted(matrix), start=4):
        put(mx, i, 1, item, text=True)
        for j, a in enumerate(aspects, start=2):
            cell = matrix[item].get(a)
            put(mx, i, j, cell["avg"] if cell else None, fmt="0.00")
    r2 = 6 + len(matrix)
    put(mx, r2 - 1, 1, "Mentions behind each cell", bold=True)
    _header(mx, r2, ["Item"] + aspects)
    for i, item in enumerate(sorted(matrix), start=r2 + 1):
        put(mx, i, 1, item, text=True)
        for j, a in enumerate(aspects, start=2):
            cell = matrix[item].get(a)
            put(mx, i, j, cell["n"] if cell else None)
    mx.column_dimensions["A"].width = 28
    if matrix:
        mx.conditional_formatting.add(f"B4:{chr(65 + len(aspects))}{3 + len(matrix)}", ColorScaleRule(
            start_type="num", start_value=-1, start_color="F8696B", mid_type="num", mid_value=0,
            mid_color="FFFFFF", end_type="num", end_value=1, end_color="63BE7B"))

    notes = ["Sentiment is on a -1..1 scale. 'rating' basis means the sentence had no sentiment words and "
             "inherited the review's star rating."]
    if join:
        sj = wb.create_sheet("Sales Join")
        put(sj, 1, 1, f'Item summary ({join["level"]} grain). units_index = item units / mean units of all items in the same state+period', bold=True)
        _header(sj, 3, ["Item", "Joined cells", "Mentions", "Avg sentiment", "Mean units index", "Sentiment vs units (r)", "Confidence"])
        for i, it in enumerate(join["items"], start=4):
            for col, k in enumerate(["item", "joined_cells", "mentions", "avg_sentiment", "mean_units_index",
                                     "sentiment_vs_units_r", "confidence"], 1):
                put(sj, i, col, it[k], text=True)
        r0 = 6 + len(join["items"])
        cols = list(join["joined_rows"][0]) if join["joined_rows"] else ["item", "region", "period"]
        put(sj, r0 - 1, 1, "Joined rows", bold=True)
        _header(sj, r0, cols)
        for i, row in enumerate(join["joined_rows"], start=r0 + 1):
            for col, k in enumerate(cols, 1):
                put(sj, i, col, row[k], text=True)
        for col in "ABCDEFGHIJ":
            sj.column_dimensions[col].width = 18
        dr = wb.create_sheet("Drivers")
        put(dr, 1, 1, "Candidate drivers for under-indexing items (hypotheses, not proof)", bold=True)
        _header(dr, 3, ["Item", "Mean units index", "Aspect", "Mentions", "Avg sentiment", "Gap vs all items", "Worst quotes"])
        r = 4
        for d in join["drivers"]:
            for g in d["candidate_drivers"] or [None]:
                put(dr, r, 1, d["item"], text=True); put(dr, r, 2, d["mean_units_index"])
                if g:
                    put(dr, r, 3, g["aspect"]); put(dr, r, 4, g["mentions"]); put(dr, r, 5, g["avg_sentiment"])
                    put(dr, r, 6, g["gap_vs_all_items"]); put(dr, r, 7, " | ".join(g["worst_quotes"]), text=True)
                r += 1
        for col, w in zip("ABCDEFG", [28, 16, 16, 10, 14, 16, 90]):
            dr.column_dimensions[col].width = w
        if join["unmatched_sales_items"]:
            notes.append("Sales items not matched to the menu catalogue (excluded): " + ", ".join(join["unmatched_sales_items"]))
        notes += join["caveats"]
    nt = wb.create_sheet("Notes")
    for i, n in enumerate(notes, start=1):
        put(nt, i, 1, n, text=True)
    nt.column_dimensions["A"].width = 140
    wb.save(path)
    return path



def build_study_insights_workbook(path: Path, study: dict, insights_data: dict, all_mentions: dict, all_reviews: dict, all_places: dict) -> Path:
    wb = Workbook()
    
    # 1. Reviews Tab
    rev_ws = wb.active
    rev_ws.title = "Reviews (Raw)"
    put(rev_ws, 1, 1, f"Extracted Reviews for Study: {study['name']}", bold=True, size=14)
    rev_cols = ["Brand", "Source", "Place ID", "Place Name", "Review ID", "Date", "Rating", "Sentiment", "Text", "URL"]
    _header(rev_ws, 3, rev_cols)
    r_idx = 4
    for brand, revs in all_reviews.items():
        for r in revs:
            put(rev_ws, r_idx, 1, brand)
            put(rev_ws, r_idx, 2, r.get("source", "N/A"), text=True)
            put(rev_ws, r_idx, 3, r.get("place_id"), text=True)
            put(rev_ws, r_idx, 4, r.get("place_name") or brand, text=True)
            put(rev_ws, r_idx, 5, r.get("review_id"), text=True)
            put(rev_ws, r_idx, 6, r.get("review_date"), text=True)
            put(rev_ws, r_idx, 7, r.get("rating"), fmt="0.0")
            put(rev_ws, r_idx, 8, sentiment_of(r.get("rating")), text=True) 
            put(rev_ws, r_idx, 9, r.get("text"), text=True)
            put(rev_ws, r_idx, 10, r.get("url"), text=True)
            r_idx += 1
    
    for col, w in zip("ABCDEFGHIJ", [15, 15, 15, 20, 15, 12, 8, 12, 80, 20]):
        rev_ws.column_dimensions[col].width = w
        
    # 2. Places Tab
    pl_ws = wb.create_sheet("Places (Raw)")
    put(pl_ws, 1, 1, "Places Extracted from Web API", bold=True, size=14)
    pl_cols = ["Brand", "Source", "Place ID", "Name", "Address", "Rating", "Review Count", "URL"]
    _header(pl_ws, 3, pl_cols)
    p_idx = 4
    for brand, places in all_places.items():
        for p in places:
            put(pl_ws, p_idx, 1, brand)
            put(pl_ws, p_idx, 2, p.get("source", "N/A"), text=True)
            put(pl_ws, p_idx, 3, p.get("place_id"), text=True)
            put(pl_ws, p_idx, 4, p.get("name"))
            put(pl_ws, p_idx, 5, p.get("address"), text=True)
            put(pl_ws, p_idx, 6, p.get("rating"))
            put(pl_ws, p_idx, 7, p.get("reviews_count"))
            put(pl_ws, p_idx, 8, p.get("url"), text=True)
            p_idx += 1
            
    for col, w in zip("ABCDEFGH", [15, 15, 15, 20, 30, 8, 12, 30]):
        pl_ws.column_dimensions[col].width = w

    # 3. Comparative Metrics
    ws = wb.create_sheet("Comparative Metrics")
    put(ws, 1, 1, f'Study Insights: {study["name"]}', bold=True, size=14)
    put(ws, 2, 1, "Focus Brand"); put(ws, 2, 2, study["focus_brand"], text=True)
    
    _header(ws, 4, ["Brand", "Reviews", "Mentions", "Share of Voice", "Mean Rating", "Mean Sentiment", "% Negative", "Sentiment vs Focus", "Rating vs Focus"])
    
    for i, b in enumerate(insights_data["brands"], start=5):
        put(ws, i, 1, b["brand"])
        put(ws, i, 2, b["reviews"])
        put(ws, i, 3, b["mentions"])
        put(ws, i, 4, b.get("share_of_voice", 0), fmt="0.0%")
        put(ws, i, 5, b.get("mean_rating"), fmt="0.00")
        put(ws, i, 6, b.get("mean_sentiment"), fmt="0.00")
        put(ws, i, 7, b.get("negative_share", 0), fmt="0.0%")
        
        delta = b.get("delta_vs_focus", {})
        put(ws, i, 8, delta.get("sentiment"), fmt="0.00")
        put(ws, i, 9, delta.get("rating"), fmt="0.00")
        
    for col, w in zip("ABCDEFGHI", [20, 10, 10, 15, 15, 15, 15, 20, 20]):
        ws.column_dimensions[col].width = w
        
    if len(insights_data["brands"]) > 0:
        ws.conditional_formatting.add(f"F5:F{4+len(insights_data['brands'])}", ColorScaleRule(
            start_type="num", start_value=-1, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=1, end_color="63BE7B"))

    # 4. Aspect Deltas
    aspects = sorted({m["aspect"] for ments in all_mentions.values() for m in ments})
    ad = wb.create_sheet("Aspect Deltas")
    put(ad, 1, 1, "Average Sentiment by Aspect & Brand", bold=True)
    _header(ad, 3, ["Aspect"] + [b["brand"] for b in insights_data["brands"]])
    
    aspect_scores = {b["brand"]: {a: [] for a in aspects} for b in insights_data["brands"]}
    for brand, ments in all_mentions.items():
        if brand in aspect_scores:
            for m in ments:
                asp = m["aspect"]
                if asp in aspect_scores[brand]:  # guard unknown aspects (e.g. 'general')
                    aspect_scores[brand][asp].append(m["sentiment"])
                
    for i, a in enumerate(aspects, start=4):
        put(ad, i, 1, a)
        for j, b in enumerate(insights_data["brands"], start=2):
            scores = aspect_scores[b["brand"]].get(a, [])
            mean_s = round(sum(scores)/len(scores), 3) if scores else None
            put(ad, i, j, mean_s, fmt="0.00")
            
    for col in "ABCDEFGHIJKLMNOPQRSTUVWXYZ"[:len(insights_data["brands"])+1]:
        ad.column_dimensions[col].width = 20
        
    # Heatmap for Aspects — last data col is B + (n_brands - 1) = chr(66 + n_brands - 1) = chr(65 + n_brands)
    if aspects and insights_data["brands"]:
        n_b = len(insights_data["brands"])
        last_col = chr(65 + n_b) if n_b <= 24 else "Z"  # safe cap at Z
        ad.conditional_formatting.add(f"B4:{last_col}{3+len(aspects)}", ColorScaleRule(
            start_type="num", start_value=-1, start_color="F8696B", mid_type="num", mid_value=0,
            mid_color="FFFFFF", end_type="num", end_value=1, end_color="63BE7B"))

    # 5. Monthly Trends
    mt = wb.create_sheet("Monthly Trends")
    put(mt, 1, 1, "Monthly Average Sentiment by Brand", bold=True)
    
    all_months = set()
    for b in insights_data["brands"]:
        all_months.update(b.get("monthly_trend", {}).keys())
    all_months = sorted(list(all_months))
    
    _header(mt, 3, ["Month"] + [b["brand"] for b in insights_data["brands"]])
    for i, mth in enumerate(all_months, start=4):
        put(mt, i, 1, mth)
        for j, b in enumerate(insights_data["brands"], start=2):
            put(mt, i, j, b.get("monthly_trend", {}).get(mth), fmt="0.00")
            
    for col in "ABCDEFGHIJKLMNOPQRSTUVWXYZ"[:len(insights_data["brands"])+1]:
        mt.column_dimensions[col].width = 20
        
    # Add Chart to Trends
    if all_months and insights_data["brands"]:
        chart = LineChart()
        chart.title = "Sentiment Trends Over Time"
        chart.y_axis.title = "Average Sentiment (-1 to 1)"
        chart.x_axis.title = "Month"
        data = Reference(mt, min_col=2, min_row=3, max_col=1+len(insights_data["brands"]), max_row=3+len(all_months))
        cats = Reference(mt, min_col=1, min_row=4, max_row=3+len(all_months))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 18
        chart.height = 10
        mt.add_chart(chart, f"{chr(67+len(insights_data['brands']))}4")

    wb.save(path)
    return path

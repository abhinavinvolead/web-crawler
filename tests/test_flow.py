"""End-to-end test in MOCK mode: no network, no credits."""
import os
import tempfile
from tests.conftest import tmp

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402
from app.normalize import is_brand_place, normalize_review  # noqa: E402


def test_full_run_and_exports():
    with TestClient(app) as c:
        assert "MOCK" in c.get("/api/v1/health").json()["mode"]
        r = c.post("/api/v1/runs", json={"sources": ["google_maps", "tripadvisor"], "max_places": 4,
                                         "tripadvisor_urls": ["https://www.tripadvisor.com/Restaurant_Review-mock"]})
        assert r.status_code == 202
        rid = r.json()["run_id"]
        run = c.get(f"/api/v1/runs/{rid}").json()
        assert run["status"] == "completed", run
        places = c.get(f"/api/v1/runs/{rid}/places").json()
        assert len(places) == 4
        assert all("Station" not in p["name"] for p in places)  # transit decoy filtered
        by = {s["source"]: s["reviews"] for s in run["by_source"]}
        assert by["google_maps"] > 20 and by["tripadvisor"] > 5
        ids = [x["review_id"] for x in c.get(f"/api/v1/runs/{rid}/reviews?limit=1000").json()["results"]]
        assert len(ids) == len(set(ids))  # duplicates removed
        assert any(j["error_records"] > 0 for j in run["jobs"])  # error records skipped, not stored
        x = c.get(f"/api/v1/runs/{rid}/export.xlsx")
        assert x.status_code == 200
        open(f"{tmp}/o.xlsx", "wb").write(x.content)
        wb = load_workbook(f"{tmp}/o.xlsx")
        assert wb.sheetnames == ["Summary", "Reviews", "Places", "Run Info"]
        assert wb["Reviews"].max_row - 1 == sum(by.values())
        assert c.get(f"/api/v1/runs/{rid}/export.csv").status_code == 200


def test_guards():
    with TestClient(app) as c:
        assert c.post("/api/v1/runs", json={"days_limit": None}).status_code == 400
        assert c.post("/api/v1/runs", json={"max_places": 500}).status_code == 400
        r = c.post("/api/v1/runs", json={"sources": ["yelp"]}).json()["run_id"]  # no URLs -> skipped
        run = c.get(f"/api/v1/runs/{r}").json()
        assert run["status"] == "failed" and "yelp_urls" in run["error"]
        assert c.get("/api/v1/runs/nope").status_code == 404


def test_resume_reclaims_stale_active_run():
    from datetime import datetime, timedelta, timezone
    from app import db

    with TestClient(app) as c:
        rid = c.post("/api/v1/runs", json={"max_places": 1}).json()["run_id"]
        old = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat(timespec="seconds")
        with db.conn() as connection:
            connection.execute(
                db.runs.update().where(db.runs.c.id == rid).values(
                    status="running", updated_at=old
                )
            )
        response = c.post(f"/api/v1/runs/{rid}/resume")
        assert response.status_code == 202


def test_formula_injection_and_filters():
    rec = {"review_id": "1", "review_rating": 1, "review": "=HYPERLINK(\"http://x\")", "place_id": "p"}
    assert normalize_review("google_maps", rec)["text"].startswith("=")  # stored raw, escaped on export
    assert normalize_review("google_maps", {"error": "dead_page"}) is None
    assert not is_brand_place({"name": "Subway Station", "category": "Transit station"}, "Subway")
    assert is_brand_place({"name": "Subway", "category": "Sandwich shop"}, "Subway")


def test_zero_reviews_is_flagged_and_raw_sample_survives(monkeypatch):
    """Provider delivers only error records: run must not say 'completed', raw must be inspectable."""
    from app import pipeline

    class ErrOnly:
        async def collect(self, **kw):
            if kw["kind"] == "discover":
                return [{"place_id": "p1", "name": "Subway", "category": "Sandwich shop",
                         "url": "https://www.google.com/maps/place/Subway/@1,2,17z/x"}]
            return [{"url": kw["inputs"][0]["url"], "error": "dead_page", "error_code": "dead_page"}]

    monkeypatch.setattr(pipeline, "get_provider", lambda: ErrOnly())
    with TestClient(app) as c:
        rid = c.post("/api/v1/runs", json={"max_places": 1}).json()["run_id"]
        run = c.get(f"/api/v1/runs/{rid}").json()
        assert run["status"] == "failed" and "0 reviews" in run["error"]
        raw = c.get(f"/api/v1/runs/{rid}/raw-sample").json()
        rev = [r for r in raw if r["kind"] == "reviews"][0]
        assert rev["raw_first_records"][0]["error"] == "dead_page"
        assert "First error record" in rev["message"]


def _sales_csv(tmp_months):
    """Meatball Marinara deliberately under-indexes; the rest sell evenly."""
    import random
    rnd = random.Random(1)
    lines = ["item,region,period,units"]
    for st in ("NY", "CA", "IL", "TX", "FL"):
        for m in tmp_months:
            for item, base in (("Meatball Marinara", 400), ("Italian B.M.T.", 1500), ("Spicy Italian", 1400),
                               ("Turkey Breast", 1450), ("Cookies", 1300)):
                lines.append(f"{item},{st},{m}-01,{int(base * (0.9 + rnd.random() * 0.2))}")
    lines.append("Mystery Sub,NY,2026-09-01,10")
    return "\n".join(lines).encode()


def test_menu_sentiment_and_sales_join():
    from datetime import date, timedelta
    today = date.today()
    months = []
    y, m = today.year, today.month
    for _ in range(6):
        months.append(f"{y}-{m:02d}")
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    with TestClient(app) as c:
        rid = c.post("/api/v1/runs", json={"max_places": 8, "days_limit": 180}).json()["run_id"]
        assert c.post(f"/api/v1/runs/{rid}/insights/export.xlsx").status_code in (404, 405)
        assert c.get(f"/api/v1/runs/{rid}/insights/items").status_code == 409  # not analysed yet
        assert c.post(f"/api/v1/runs/{rid}/analyze", json={"method": "llm"}).status_code == 400  # no key
        assert c.post(f"/api/v1/runs/{rid}/analyze", json={"method": "lexicon"}).status_code == 202
        assert c.get(f"/api/v1/runs/{rid}/analysis").json()["status"] == "completed"
        items = {i["item"]: i for i in c.get(f"/api/v1/runs/{rid}/insights/items").json()}
        assert items["Meatball Marinara"]["avg_sentiment"] < 0 < items["Italian B.M.T."]["avg_sentiment"]
        up = c.post("/api/v1/sales/upload", files={"file": ("s.csv", _sales_csv(months), "text/csv")}).json()
        assert up["rows"] > 100 and up["unmatched_items"] == ["Mystery Sub"]
        j = c.get(f"/api/v1/runs/{rid}/insights/sales-join", params={"sales_id": up["sales_id"]}).json()
        assert j["joined_rows"] and all("revenue" not in r for r in j["joined_rows"])  # finance-only by default
        mm = [i for i in j["items"] if i["item"] == "Meatball Marinara"][0]
        assert mm["mean_units_index"] < 0.6
        drv = [d for d in j["drivers"] if d["item"] == "Meatball Marinara"]
        assert drv and drv[0]["candidate_drivers"]
        x = c.get(f"/api/v1/runs/{rid}/insights/export.xlsx", params={"sales_id": up["sales_id"]})
        assert x.status_code == 200
        open(f"{tmp}/i.xlsx", "wb").write(x.content)
        assert {"Items", "Aspect Matrix", "Sales Join", "Drivers", "Notes"} <= set(load_workbook(f"{tmp}/i.xlsx").sheetnames)
        bad = c.post("/api/v1/sales/upload", files={"file": ("b.csv", b"a,b\n1,2", "text/csv")})
        assert bad.status_code == 400


def test_llm_analyzer_with_stubbed_api(monkeypatch):
    import asyncio
    import json as _json

    import httpx

    from app import analyzers

    def handler(req):
        body = _json.loads(req.content)
        ids = [ln.split("]")[0][1:] for ln in body["messages"][0]["content"].splitlines() if ln.startswith("[")]
        out = [{"id": i, "mentions": [{"item": "meatball", "aspect": "taste", "sentiment": -0.8, "quote": "bland"},
                                      {"item": "Cookies", "aspect": "nope", "sentiment": "1.5", "quote": "great"}]}
               for i in ids]
        return httpx.Response(200, json={"content": [{"type": "text", "text": _json.dumps(out)}]})

    real = httpx.AsyncClient
    monkeypatch.setattr(analyzers.httpx, "AsyncClient", lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    menu = {"Meatball Marinara": ["meatball"], "Cookies": []}
    revs = [{"review_id": f"r{i}", "source": "google_maps", "rating": 2, "text": "x"} for i in range(20)]
    out = asyncio.run(analyzers.LLMAnalyzer("k", "m").analyze(revs, menu))
    assert len(out) == 40
    assert {m["item"] for m in out} == {"Meatball Marinara", "Cookies"}  # alias resolved to canonical
    assert all(m["aspect"] in ("taste", "general") and -1 <= m["sentiment"] <= 1 for m in out)  # clamped/validated


def test_no_peers_error_retry_and_fallback(monkeypatch):
    """Reproduces the live failure: review scraper returns 502 no_peers for every URL."""
    from app import pipeline
    state = {"calls": 0}
    place = {"place_id": "p1", "name": "Subway", "category": "Sandwich shop", "address": "1 A St, Chicago, IL 60612",
             "url": "https://www.google.com/maps/place/Subway/data=!x",
             "top_reviews": [{"reviewer_name": "A", "review_date": "2099-01-01T00:00:00Z", "rating": 1,
                              "content": "The bathroom was dirty and the bread was stale."}]}

    class Flaky:
        async def collect(self, **kw):
            if kw["kind"] == "discover":
                return [place]
            state["calls"] += 1
            if state["calls"] == 1:  # first attempt: platform-side error, echoed under "input"
                return [{"input": {"url": kw["inputs"][0]["url"]}, "error": "Crawler error: Unexpected Status 502 (no_peer)",
                         "error_code": "no_peers"}]
            assert [i["url"] for i in kw["inputs"]] == [place["url"]]  # only the failed URL is retried
            return [{"review_id": "real-1", "review": "Fresh and fast.", "review_rating": 5,
                     "review_date": "2099-01-02", "place_id": "p1", "url": place["url"]}]

    monkeypatch.setattr(pipeline, "get_provider", lambda: Flaky())
    with TestClient(app) as c:
        rid = c.post("/api/v1/runs", json={"max_places": 1, "days_limit": 3650}).json()["run_id"]
        run = c.get(f"/api/v1/runs/{rid}").json()
        assert run["status"] == "completed_with_warnings", run
        job = [j for j in run["jobs"] if j["kind"] == "reviews"][0]
        assert "FALLBACK" in job["message"] and "retry_failed=true" in job["message"]
        ids = [r["review_id"] for r in c.get(f"/api/v1/runs/{rid}/reviews").json()["results"]]
        assert ids and all(i.startswith("top-") for i in ids)
        assert c.post(f"/api/v1/runs/{rid}/resume").status_code == 202  # plain resume must not re-bill
        assert state["calls"] == 1
        assert c.post(f"/api/v1/runs/{rid}/resume?retry_failed=true").status_code == 202
        assert state["calls"] == 2
        ids = [r["review_id"] for r in c.get(f"/api/v1/runs/{rid}/reviews").json()["results"]]
        assert ids == ["real-1"]  # real review supersedes the fallback
        assert c.get(f"/api/v1/runs/{rid}").json()["status"] == "completed"
        assert c.post(f"/api/v1/runs/{rid}/resume?retry_failed=true").status_code == 409  # nothing left


def test_snapshot_timeout_keeps_embedded_reviews(monkeypatch):
    from app import pipeline
    from app.brightdata import BrightDataTransientError

    class TimedOut:
        async def collect(self, **kw):
            if kw["kind"] == "discover":
                return [{
                    "place_id": "p-timeout", "name": "Subway", "category": "Sandwich shop",
                    "url": "https://www.google.com/maps/place/Subway/@1,2,17z/x",
                    "top_reviews": [{
                        "reviewer_name": "A", "review_date": "2099-01-01",
                        "rating": 4, "content": "Fresh and fast."
                    }]
                }]
            raise BrightDataTransientError("snapshot still running")

    monkeypatch.setattr(pipeline, "get_provider", lambda: TimedOut())
    with TestClient(app) as c:
        rid = c.post("/api/v1/runs", json={"max_places": 1}).json()["run_id"]
        run = c.get(f"/api/v1/runs/{rid}").json()
        assert run["status"] == "completed_with_warnings"
        assert c.get(f"/api/v1/runs/{rid}/reviews").json()["count"] == 1


def test_snapshot_timeout_keeps_past_dated_embedded_reviews(monkeypatch):
    """Reproduces the real-world bug where review dates are from months ago and days_limit=30."""
    from app import pipeline
    from app.brightdata import BrightDataTransientError

    class TimedOutWithPastReviews:
        async def collect(self, **kw):
            if kw["kind"] == "discover":
                return [{
                    "place_id": "p-past", "name": "Burger King", "category": "Fast food restaurant",
                    "url": "https://www.google.com/maps/place/Burger+King/@1,2,17z/x",
                    "top_reviews": [{
                        "reviewer_name": "Elle G.", "review_date": "2024-01-15T00:00:00Z",
                        "rating": 5, "content": "Great burgers and friendly service."
                    }]
                }]
            raise BrightDataTransientError("snapshot still running after 300s")

    monkeypatch.setattr(pipeline, "get_provider", lambda: TimedOutWithPastReviews())
    with TestClient(app) as c:
        rid = c.post("/api/v1/runs", json={"brand": "Burger King", "max_places": 1, "days_limit": 30}).json()["run_id"]
        run = c.get(f"/api/v1/runs/{rid}").json()
        assert run["status"] == "completed_with_warnings"
        revs = c.get(f"/api/v1/runs/{rid}/reviews").json()
        assert revs["count"] == 1
        assert "Great burgers" in revs["results"][0]["text"]


def test_build_review_inputs_sets_limit():
    from app.registry import get_source
    adapter = get_source("google_maps")
    inputs = adapter.build_review_inputs(["https://google.com/maps/place/1"], days_limit=30, max_reviews=25)
    assert len(inputs) == 1
    assert "limit" not in inputs[0]  # Bright Data gd_luzfs1dn2oa0teb81 rejects limit with 400
    assert inputs[0]["days_limit"] == 30


def test_create_study_flow():
    with TestClient(app) as c:
        # Test /api/v1/brands endpoint
        brands_res = c.get("/api/v1/brands")
        assert brands_res.status_code == 200
        brands_list = brands_res.json()
        assert "Subway" in brands_list and "McDonald's" in brands_list

        payload = {
            "name": "Subway vs McD",
            "focus_brand": "Subway",
            "competitors": ["McDonald's"],
            "sources": ["google_maps"],
            "days_limit": 30,
            "max_places": 2,
            "study_budget": 500
        }
        res = c.post("/api/v1/studies", json=payload)
        assert res.status_code == 202
        data = res.json()
        assert "study_id" in data and len(data["run_ids"]) == 2
        sid = data["study_id"]
        study = c.get(f"/api/v1/studies/{sid}").json()
        assert len(study["runs"]) == 2
        for run in study["runs"]:
            assert run["status"] in ("completed", "completed_with_warnings")

        # Run analysis on all runs in the study
        for rid in data["run_ids"]:
            an_res = c.post(f"/api/v1/runs/{rid}/analyze", json={"method": "lexicon"})
            assert an_res.status_code == 202
            assert c.get(f"/api/v1/runs/{rid}/analysis").json()["status"] == "completed"

        # Verify study insights
        insights_res = c.get(f"/api/v1/studies/{sid}/insights")
        assert insights_res.status_code == 200
        ins = insights_res.json()
        assert ins["focus_brand"] == "Subway"
        assert len(ins["brands"]) == 2

        # Verify study excel export
        exp_res = c.get(f"/api/v1/studies/{sid}/export.xlsx")
        assert exp_res.status_code == 200
        assert len(exp_res.content) > 1000

        # Verify per-brand reviews CSV export
        for rid in data["run_ids"]:
            csv_res = c.get(f"/api/v1/runs/{rid}/export.csv")
            assert csv_res.status_code == 200
            assert b"review_id" in csv_res.content


def test_raw_sample_corrupted_json_handled_safely():
    from app import db
    rid = db.create_run("Subway", {"brand": "Subway"})
    # Insert a job with a deliberately corrupted/unterminated JSON sample string
    db.upsert_job(rid, "google_maps", "reviews", status="done", records=5,
                  sample='[{"url": "https://example.com", "text": "Unfinished string')
    with TestClient(app) as c:
        res = c.get(f"/api/v1/runs/{rid}/raw-sample")
        assert res.status_code == 200
        data = res.json()
        assert len(data) == 1
        assert "Unfinished string" in data[0]["raw_first_records"][0]["raw_preview"]


def test_burger_king_study_and_excel_export():
    import io
    from openpyxl import load_workbook
    with TestClient(app) as c:
        payload = {
            "name": "BK Test Study",
            "focus_brand": "Burger King",
            "competitors": [],
            "sources": ["google_maps"],
            "days_limit": 30,
            "max_places": 2,
            "study_budget": 500,
        }
        res = c.post("/api/v1/studies", json=payload)
        assert res.status_code == 202
        sid = res.json()["study_id"]
        rid = res.json()["run_ids"][0]

        # Verify run reviews are for Burger King
        revs = c.get(f"/api/v1/runs/{rid}/reviews?limit=100").json()["results"]
        assert len(revs) > 0
        assert all(r["place_name"] == "Burger King" for r in revs)
        assert all(r["brand"] == "Burger King" for r in revs)
        # Ensure text is not Subway
        assert not any("meatball marinara" in r["text"].lower() for r in revs)
        assert not any("italian bmt" in r["text"].lower() for r in revs)

        # Verify Excel export
        exp_res = c.get(f"/api/v1/studies/{sid}/export.xlsx")
        assert exp_res.status_code == 200
        wb = load_workbook(io.BytesIO(exp_res.content))
        rev_sheet = wb["Reviews (Raw)"]
        assert rev_sheet.max_row > 3
        # Check first data row (row 4)
        brand_cell = rev_sheet.cell(4, 1).value
        place_cell = rev_sheet.cell(4, 4).value
        sentiment_cell = rev_sheet.cell(4, 8).value
        assert brand_cell == "Burger King"
        assert place_cell == "Burger King"
        assert sentiment_cell in ("Positive", "Neutral", "Negative")




